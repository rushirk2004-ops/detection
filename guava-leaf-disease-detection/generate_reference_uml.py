import sys
import requests
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import time

diagrams = {
    "Use Case Diagram": """@startuml
skinparam handwritten false
skinparam monochrome false
skinparam usecase {
  BackgroundColor #cce5ff
  BorderColor #005cbf
}
left to right direction
actor User
actor Admin

rectangle "Guava Fruit Disease Detection System" {
  usecase "Upload Fruit Image" as UC1
  usecase "View Preprocessed Image" as UC2
  usecase "Get Disease Prediction" as UC3
  usecase "View Disease Details" as UC4
  usecase "Manage Dataset (Admin)" as UC5
  usecase "Logout" as UC6
}

User --> UC1
User --> UC2
User --> UC3
User --> UC4
User --> UC6
Admin --> UC5
Admin --> UC6
@enduml""",

    "Class Diagram": """@startuml
skinparam class {
  BackgroundColor #f2f2f2
  BorderColor black
}
class User << (C,#e2f0d9) >> {
  - userId : int
  - name : string
  - email : string
  - password : string
  + login()
  + uploadImage()
  + viewResult()
}
class Image << (C,#f2e6ff) >> {
  - imageId : int
  - filePath : string
  - uploadDate : Date
  - status : string
  + preprocess()
}
class Preprocessing << (C,#e6f7ff) >> {
  - preprocessId : int
  - method : string
  + resize()
  + normalize()
  + augment()
}
class DiseaseClassifier << (C,#ffe6e6) >> {
  - modelName : string
  - accuracy : float
  + extractFeatures()
  + classify()
  + predict()
}
class Prediction << (C,#fff2cc) >> {
  - predictionId : int
  - diseaseName : string
  - confidence : float
  - date : Date
  + getResult()
}
class Database << (C,#ffffe6) >> {
  - dbId : int
  - name : string
  + storeImage()
  + storeResult()
  + getHistory()
}

User "1" -right- "0..*" Image : uploads >
Image "1" -down- "1" Preprocessing : processed by >
DiseaseClassifier "1" -right- "1" Preprocessing : feeds >
DiseaseClassifier "1" -down- "1" Prediction : predicts >
Prediction "1" -right- "1" Database : stores >
@enduml""",

    "Activity Diagram": """@startuml
skinparam activity {
  BackgroundColor #cce5ff
  BorderColor #005cbf
}
start
:Open Application;
repeat
  :Upload Guava Fruit Image;
backward: Invalid Image;
repeat while (Valid Image?) is (No) not (Yes)
:Preprocess Image\n(resize, normalize, augment);
:Extract Features (CNN Model);
:Classify Disease\n(Trained Model);
:Display Result\n(Disease Name & Confidence);
if (Another Image?) then (Yes)
  -> Upload Guava Fruit Image;
else (No)
endif
stop
@enduml""",

    "Sequence Diagram": """@startuml
skinparam participant {
  BackgroundColor #e6f7ff
  BorderColor #005cbf
}
actor User
participant "Web / Mobile\\nApp (UI)" as UI
participant "Preprocessing\\nModule" as Prep << (C,#e2f0d9) >>
participant "CNN Model" as Model << (C,#f2e6ff) >>
participant "Database" as DB << (C,#fff2cc) >>

User -> UI : 1. Upload Image
activate UI
UI -> Prep : 2. Send Image
activate Prep
Prep -> Prep : 3. Preprocess Image
Prep -> Model : 4. Extract Features\\n& Classify
activate Model
Model --> UI : 5. Prediction Result
deactivate Model
deactivate Prep
UI -> DB : 6. Store Result
activate DB
deactivate DB
UI --> User : 7. Display Result
User -> UI : 8. View Result
deactivate UI
@enduml""",

    "Component Diagram": """@startuml
skinparam component {
  BackgroundColor #e6f7ff
  BorderColor #005cbf
}
component Frontend [
  Frontend (Web / Mobile App)
  ---
  - UI Components
  - Image Upload
  - Result Display
]
component Backend [
  Backend (API Server)
  ---
  - REST API
  - Business Logic
  - Model Integration
]
component DB [
  Database
  ---
  - Store Images
  - Store Results
  - Manage Dataset
]
component Model [
  CNN Model (Trained Model)
  ---
  - Feature Extraction
  - Disease Classification
]

Frontend -right-> Backend : HTTP/REST
Frontend -down-> DB : Fetch/Display Results
Backend -down-> Model : Model Request/Response
Backend -down-> DB : Store/Retrieve Data
@enduml""",

    "Deployment Diagram": """@startuml
skinparam node {
  BackgroundColor #e6f7ff
  BorderColor #005cbf
}
node "User Device\\n(PC / Mobile)" as Device << (C,#e6f7ff) >> {
  [Browser / App]
  [User Interface]
}
node "Web Server\\n(Cloud / Local)" as WebServer << (C,#f2e6ff) >> {
  [Frontend Files]
  [API Gateway]
}
node "Application Server\\n(Backend)" as AppServer << (C,#e2f0d9) >> {
  [Business Logic]
  [CNN Model]
  [REST API]
}
node "Database Server\\n(Cloud / On-Premise)" as DBServer << (C,#ffe6e6) >> {
  [Store Images]
  [Store Results]
  [Manage Dataset]
}

Device -right- WebServer : HTTPS
WebServer -down- AppServer : Internal Network
AppServer -left- DBServer : Connection
@enduml"""
}

def generate_image_from_kroki(name, puml_text):
    url = 'https://kroki.io/plantuml/png'
    headers = {'Content-Type': 'text/plain'}
    response = requests.post(url, headers=headers, data=puml_text.encode('utf-8'))
    if response.status_code == 200:
        filename = f"{name.replace(' ', '_').lower()}.png"
        with open(filename, 'wb') as f:
            f.write(response.content)
        return filename
    else:
        print(f"Failed to generate {name}. Status code: {response.status_code}")
        print(f"Response: {response.text}")
        return None

def create_doc(output_path):
    doc = Document()
    
    title = doc.add_heading('Guava Fruit Disease Detection', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = doc.add_paragraph('UML Reference Diagrams')
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].bold = True
    subtitle.runs[0].font.size = Pt(16)
    
    desc = doc.add_paragraph('This document contains the UML diagrams strictly modeled after the provided reference architecture, including all specified components (Database, API Server, Mobile App, etc.). Each diagram is rendered in high resolution on its own page.')
    desc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    desc.runs[0].italic = True
    
    doc.add_page_break()

    for idx, (name, puml) in enumerate(diagrams.items(), 1):
        print(f"Generating image for {name}...")
        img_path = generate_image_from_kroki(name, puml)
        time.sleep(1)
        
        if img_path:
            h = doc.add_heading(f'{idx}. {name}', level=1)
            try:
                pic = doc.add_picture(img_path, width=Inches(6.0))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            except Exception as e:
                print(f"Error loading image {img_path}: {e}")
            
            if idx < len(diagrams):
                doc.add_page_break()
                
    doc.save(output_path)
    print(f"UML Document successfully created at {output_path}")

if __name__ == '__main__':
    out_path = sys.argv[1]
    create_doc(out_path)
