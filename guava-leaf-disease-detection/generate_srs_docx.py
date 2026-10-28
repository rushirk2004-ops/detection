import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def add_heading(doc, text, level):
    h = doc.add_heading(text, level=level)
    if level == 1:
        h.runs[0].font.color.rgb = RGBColor(0x00, 0x33, 0x66)

def add_paragraph(doc, text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    return p

def create_srs(output_path):
    doc = Document()
    
    # Title Page
    title = doc.add_heading('Software Requirements Specification (SRS)', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    subtitle = doc.add_paragraph('Guava Leaf & Fruit Disease Detection System')
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].bold = True
    subtitle.runs[0].font.size = Pt(24)
    
    doc.add_paragraph('\n\n\n\n\n')
    doc.add_paragraph('Prepared by: Development Team\nDate: October 2026\nVersion: 1.0').alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()

    # Table of Contents placeholder
    add_heading(doc, 'Table of Contents', 1)
    doc.add_paragraph('1. Introduction\n2. Overall Description\n3. External Interface Requirements\n4. System Features\n5. Nonfunctional Requirements\n6. Appendices')
    doc.add_page_break()

    # 1. Introduction
    add_heading(doc, '1. Introduction', 1)
    
    add_heading(doc, '1.1 Purpose', 2)
    doc.add_paragraph('The purpose of this document is to present a detailed description of the Guava Leaf and Fruit Disease Detection System. It will explain the purpose and features of the system, the interfaces of the system, what the system will do, the constraints under which it must operate, and how the system will react to external stimuli. This document is intended for both the stakeholders and the developers of the system.')

    add_heading(doc, '1.2 Document Conventions', 2)
    doc.add_paragraph('This document follows standard IEEE SRS formatting. Main sections are numbered sequentially. Bold text is used for emphasis and defining key terms. The priority of requirements is implicit unless stated otherwise (all core features are critical).')

    add_heading(doc, '1.3 Intended Audience and Reading Suggestions', 2)
    doc.add_paragraph('This document is intended for project managers, software developers, machine learning engineers, and agricultural stakeholders. Developers and engineers should focus heavily on Section 3 (Interfaces), Section 4 (System Features), and Section 5 (Nonfunctional Requirements). Project managers should read the entire document to grasp the scope.')

    add_heading(doc, '1.4 Product Scope', 2)
    doc.add_paragraph('The Guava Disease Detection System is an intelligent software solution designed to assist farmers and agricultural experts in identifying diseases in guava plants. By leveraging deep learning models (EfficientNet-B3 and ResNet-50) via a Streamlit-based web interface, users can upload images of guava leaves or fruits and receive real-time, highly accurate disease classifications. The system goes beyond simple predictions by integrating Grad-CAM (Gradient-weighted Class Activation Mapping) to provide visual explainability—highlighting the exact regions of the image that contributed to the model’s decision. This builds trust and transparency for the end-user. The scope includes image preprocessing, local or cloud-based model inference, and detailed visual reporting.')

    add_heading(doc, '1.5 References', 2)
    doc.add_paragraph('- PyTorch Documentation (Deep Learning Framework)\n- Streamlit Documentation (Web Application Framework)\n- EfficientNet and ResNet Architecture Whitepapers\n- IEEE Guide for Software Requirements Specifications (IEEE Std 830-1998)')

    doc.add_page_break()

    # 2. Overall Description
    add_heading(doc, '2. Overall Description', 1)
    
    add_heading(doc, '2.1 Product Perspective', 2)
    doc.add_paragraph('This system is a new, self-contained application. It operates as a web-based client-server system where the frontend (Streamlit) interacts seamlessly with the backend Python environment hosting the PyTorch models. The system can be deployed locally on a user\'s machine or hosted on a cloud server to provide wider accessibility via mobile and desktop browsers.')

    add_heading(doc, '2.2 Product Functions', 2)
    doc.add_paragraph('The primary functions of the system include:\n- Secure uploading of JPG/PNG images of guava leaves/fruits.\n- Automated image preprocessing (resizing to 224x224, tensor conversion, ImageNet normalization).\n- Real-time deep learning inference using pre-trained convolutional neural networks.\n- Display of predicted disease class (e.g., Canker, Rust, Dot, Mummification, Healthy) alongside confidence percentages.\n- Generation and display of Grad-CAM heatmaps overlaid on the original image to explain the prediction.\n- Historical logging of predictions into a backend database for further analysis.')

    add_heading(doc, '2.3 User Classes and Characteristics', 2)
    doc.add_paragraph('1. Farmers/End Users: Individuals with minimal technical expertise seeking quick disease diagnosis for their crops. They require an intuitive, fast, and easy-to-use interface accessible from mobile devices.\n2. Agricultural Experts: Users who require detailed confidence scores and Grad-CAM visualizations to verify the AI\'s diagnostic logic.\n3. System Administrators: IT personnel responsible for managing the system, maintaining the database, and updating model weights.')

    add_heading(doc, '2.4 Operating Environment', 2)
    doc.add_paragraph('The frontend is accessible via any modern web browser (Chrome, Firefox, Safari, Edge) on Desktop, iOS, or Android devices. The backend server requires a Python 3.9+ environment with PyTorch, Streamlit, OpenCV, and NumPy installed. Hardware acceleration (NVIDIA CUDA GPU) is highly recommended for faster inference, though CPU execution is supported as a fallback.')

    add_heading(doc, '2.5 Design and Implementation Constraints', 2)
    doc.add_paragraph('- The system must rely on Python and PyTorch for the machine learning backbone.\n- Image uploads are restricted to a maximum size (e.g., 10MB) to prevent server overload.\n- Models must accept fixed 224x224 input sizes, necessitating rigorous preprocessing.\n- The system must function efficiently even over moderate internet connections (e.g., 3G/4G networks).')

    add_heading(doc, '2.6 Assumptions and Dependencies', 2)
    doc.add_paragraph('- It is assumed that the uploaded images are reasonably clear, in-focus, and well-lit.\n- The system depends on the pre-trained EfficientNet-B3 and ResNet-50 weights being present in the local file system.\n- It is assumed that the server hosting the application has sufficient RAM to load the deep learning models into memory.')

    doc.add_page_break()

    # 3. External Interface Requirements
    add_heading(doc, '3. External Interface Requirements', 1)

    add_heading(doc, '3.1 User Interfaces', 2)
    doc.add_paragraph('The user interface will be built using Streamlit. It will feature a minimalist, responsive design. The main dashboard will consist of a sidebar for selecting the analysis target (Leaf vs. Fruit) and the underlying model architecture. The main panel will feature a drag-and-drop file upload zone. Upon submission, a progress spinner will be displayed, followed by the final results laid out in clearly defined sections: The original image, a prominent success/alert banner showing the disease name, a metric widget displaying confidence, and the Grad-CAM heatmap visualization.')

    add_heading(doc, '3.2 Hardware Interfaces', 2)
    doc.add_paragraph('No specialized hardware interfaces are required on the client side other than a standard device camera (if taking photos directly) and internet connectivity. On the server side, the system interfaces with the host CPU and GPU (via CUDA drivers) to accelerate tensor operations.')

    add_heading(doc, '3.3 Software Interfaces', 2)
    doc.add_paragraph('- PyTorch Framework: The core engine for loading `.pt` files and executing forward passes.\n- PIL/OpenCV: Utilized for decoding JPEG/PNG files and performing matrix manipulations.\n- Streamlit Framework: Acts as both the web server and the frontend UI rendering engine.\n- Local Database/Storage: Used to persist session states and log historical prediction data (e.g., SQLite or PostgreSQL).')

    add_heading(doc, '3.4 Communications Interfaces', 2)
    doc.add_paragraph('The system communicates over standard HTTP/HTTPS protocols (Ports 80/443 or default Streamlit Port 8501). Data transferred between the client and server consists primarily of binary image data (multipart/form-data) and JSON-formatted text responses.')

    doc.add_page_break()

    # 4. System Features
    add_heading(doc, '4. System Features', 1)

    # Feature 1
    add_heading(doc, '4.1 Image Upload and Validation', 2)
    doc.add_paragraph('4.1.1 Description and Priority\nHigh Priority. The user must be able to upload images for analysis safely.\n\n4.1.2 Stimulus/Response Sequences\nUser drags an image into the upload box. System validates file extension (jpg, jpeg, png) and size. If valid, the image is passed to preprocessing; if invalid, an error toast is displayed.\n\n4.1.3 Functional Requirements\n- REQ-4.1.1: The system shall accept .jpg, .jpeg, and .png formats.\n- REQ-4.1.2: The system shall reject files larger than 10MB.\n- REQ-4.1.3: The system shall gracefully handle corrupted image headers.')

    # Feature 2
    add_heading(doc, '4.2 Image Preprocessing Pipeline', 2)
    doc.add_paragraph('4.2.1 Description and Priority\nHigh Priority. Raw images cannot be fed to CNNs; they must be standardized.\n\n4.2.2 Stimulus/Response Sequences\nSystem receives raw PIL image. It resizes to 224x224, converts to PyTorch tensor, applies ImageNet normalization (Mean: [0.485, 0.456, 0.406], Std: [0.229, 0.224, 0.225]), and adds a batch dimension.\n\n4.2.3 Functional Requirements\n- REQ-4.2.1: Image must be resized exactly to 224x224 pixels without severe aspect ratio distortion (padding if necessary).\n- REQ-4.2.2: Tensor output must strictly match the data type expected by the model (Float32).')

    # Feature 3
    add_heading(doc, '4.3 Deep Learning Inference', 2)
    doc.add_paragraph('4.3.1 Description and Priority\nCritical Priority. This is the core intelligence of the application.\n\n4.3.2 Stimulus/Response Sequences\nUser triggers prediction. Preprocessed tensor is sent to the selected model (EfficientNet-B3 or ResNet-50). The model performs a forward pass under torch.no_grad() and outputs raw logits.\n\n4.3.3 Functional Requirements\n- REQ-4.3.1: The system shall cache the loaded model in memory to prevent reloading on every request.\n- REQ-4.3.2: The system shall correctly map output logits to the 5 designated classes (canker, dot, healthy, mummification, rust).')

    # Feature 4
    add_heading(doc, '4.4 Confidence Calculation and Softmax', 2)
    doc.add_paragraph('4.4.1 Description and Priority\nHigh Priority. Logits must be converted to understandable probabilities.\n\n4.4.2 Stimulus/Response Sequences\nLogits are passed through a Softmax activation layer. The highest probability is extracted and multiplied by 100 to yield a confidence percentage.\n\n4.4.3 Functional Requirements\n- REQ-4.4.1: Confidence scores must be displayed to two decimal places (e.g., 98.45%).')

    # Feature 5
    add_heading(doc, '4.5 Grad-CAM Explainability', 2)
    doc.add_paragraph('4.5.1 Description and Priority\nMedium-High Priority. Enhances user trust by visualizing AI focus areas.\n\n4.5.2 Stimulus/Response Sequences\nSystem triggers backward pass to extract gradients from the final convolutional layer (e.g., layer4 in ResNet). It generates a weighted heatmap, resizes it to the original image dimensions, applies a colormap (e.g., JET), and blends it with the original image.\n\n4.5.3 Functional Requirements\n- REQ-4.5.1: The heatmap must accurately align with the structural features of the original image.\n- REQ-4.5.2: The visualization must clearly highlight regions of high activation in red/yellow and low activation in blue.')

    doc.add_page_break()

    # 5. Nonfunctional Requirements
    add_heading(doc, '5. Other Nonfunctional Requirements', 1)

    add_heading(doc, '5.1 Performance Requirements', 2)
    doc.add_paragraph('- Total time from image upload to displaying the prediction and Grad-CAM result shall not exceed 3 seconds when running on a CUDA-enabled GPU.\n- On CPU environments, total execution time shall not exceed 10 seconds.\n- The system must support at least 50 concurrent users without severe degradation in response time (assuming proper backend scaling).')

    add_heading(doc, '5.2 Safety Requirements', 2)
    doc.add_paragraph('- The system shall present a disclaimer stating that the AI prediction is for informational purposes only and should not replace professional agricultural lab testing.\n- Fallback mechanisms must be in place so that if the inference engine crashes, the web interface remains active and displays a friendly error message.')

    add_heading(doc, '5.3 Security Requirements', 2)
    doc.add_paragraph('- Uploaded images must be temporarily stored in volatile memory or securely deleted from the disk immediately after inference to protect user privacy.\n- If deployed to the web, the server must force HTTPS connections to encrypt image data in transit.\n- Admin panels (if implemented) must require secure authentication and role-based access control (RBAC).')

    add_heading(doc, '5.4 Software Quality Attributes', 2)
    doc.add_paragraph('- Reliability: The system should have an uptime of 99.9%.\n- Maintainability: The codebase shall be fully documented following PEP-8 standards, and deep learning models must be easily swappable without altering the frontend logic.\n- Usability: The interface must be completely intuitive, requiring zero training for a farmer to successfully upload an image and read the result.')

    doc.add_page_break()

    # 6. Appendices
    add_heading(doc, '6. Appendices', 1)
    
    add_heading(doc, 'Appendix A: Glossary', 2)
    doc.add_paragraph('- CNN: Convolutional Neural Network. A type of deep learning model specialized for image processing.\n- Grad-CAM: Gradient-weighted Class Activation Mapping. A technique for producing visual explanations for decisions from a large class of CNNs.\n- Logit: The raw, unnormalized scores output by the final layer of a neural network before softmax is applied.\n- Tensor: A multi-dimensional array used as the fundamental data structure in PyTorch.\n- Streamlit: An open-source Python framework for building custom web apps for machine learning and data science.')

    add_heading(doc, 'Appendix B: To Be Determined (TBD) List', 2)
    doc.add_paragraph('- Expansion of dataset to include additional fruit diseases (currently under construction in UI).\n- Integration of a cloud-hosted relational database (e.g., AWS RDS) for tracking user history.\n- Development of a native offline mobile application (Android/iOS) for field workers without internet access.')
    
    doc.save(output_path)
    print(f"SRS Document successfully created at {output_path}")

if __name__ == '__main__':
    out_path = sys.argv[1]
    create_srs(out_path)
