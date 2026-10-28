import sys
from docx import Document
from docx.shared import Inches

def create_doc(image_path, output_path):
    doc = Document()
    doc.add_heading('System Architecture: Guava Leaf Disease Detection', 0)

    doc.add_paragraph('This document outlines the step-by-step deep system architecture for the Guava Leaf Disease Detection application. The architecture is designed to handle user requests, process images, predict disease presence using a deep learning model, and return actionable results.')

    doc.add_heading('System Architecture Diagram', level=1)
    doc.add_paragraph('The following diagram illustrates the complete data flow and components involved in the system.')
    
    try:
        doc.add_picture(image_path, width=Inches(6.0))
    except Exception as e:
        print(f"Error adding picture: {e}")
        
    doc.add_heading('Step-by-Step Architecture Explanation', level=1)

    # Step 1
    doc.add_heading('1. Client Side (User Interface)', level=2)
    doc.add_paragraph('The client side is the entry point for the user. It can be a web interface or a mobile application (Android/iOS).')
    p1 = doc.add_paragraph()
    p1.add_run('Step-by-step flow:\n').bold = True
    p1.add_run('- The user opens the application and selects an image of a guava leaf suspected of having a disease.\n')
    p1.add_run('- The interface compresses the image slightly for efficient transmission if necessary.\n')
    p1.add_run('- An HTTP POST request is triggered, sending the image payload to the Backend API server.')

    # Step 2
    doc.add_heading('2. Backend System (API Server & Inference Engine)', level=2)
    doc.add_paragraph('The backend acts as a bridge between the client application and the deep learning model. It is typically built using RESTful frameworks like Flask or FastAPI.')
    p2 = doc.add_paragraph()
    p2.add_run('Step-by-step flow:\n').bold = True
    p2.add_run('- The API receives the image file, validates the file format, and securely saves it temporarily.\n')
    p2.add_run('- The request is passed to the Inference Engine, which loads the pre-trained Deep Learning Model.\n')
    p2.add_run('- The Backend may log the request and user details into a Database for history and analytics.\n')
    p2.add_run('- After prediction, the Backend formulates a JSON response containing the predicted disease, confidence score, and recommended actions.')

    # Step 3
    doc.add_heading('3. Data Preprocessing Pipeline (Model Input Preparation)', level=2)
    doc.add_paragraph('Before the model can process the image, it must be transformed into the correct numerical format.')
    p3 = doc.add_paragraph()
    p3.add_run('Step-by-step flow:\n').bold = True
    p3.add_run('- Resizing: The image is resized to match the input dimensions expected by the model (e.g., 224x224 pixels).\n')
    p3.add_run('- Normalization: Pixel values (0-255) are scaled down (e.g., 0 to 1) to help the neural network converge quickly.\n')
    p3.add_run('- Tensor Conversion: The image matrix is converted into a multi-dimensional array (Tensor) ready for the neural network.')

    # Step 4
    doc.add_heading('4. Deep Learning Model (CNN Feature Extraction & Classification)', level=2)
    doc.add_paragraph('This is the core intelligence of the system, typically a Convolutional Neural Network (CNN) like ResNet or Custom CNN.')
    p4 = doc.add_paragraph()
    p4.add_run('Step-by-step flow:\n').bold = True
    p4.add_run('- Convolutional Layers: Extract low-level features (edges, colors) and high-level features (disease spots, lesions) from the leaf image.\n')
    p4.add_run('- Pooling Layers: Downsample the feature maps to reduce computational load and control overfitting.\n')
    p4.add_run('- Fully Connected Layers (Dense): Flatten the features and use learned weights to classify the image into specific disease categories (e.g., Canker, Rust, Healthy).\n')
    p4.add_run('- Softmax Output: Outputs the probability score for each class.')

    # Step 5
    doc.add_heading('5. Result Delivery and Display', level=2)
    doc.add_paragraph('Closing the loop by presenting the results to the user.')
    p5 = doc.add_paragraph()
    p5.add_run('Step-by-step flow:\n').bold = True
    p5.add_run('- The API sends the JSON response back over the network to the Client App.\n')
    p5.add_run('- The UI parses the response and dynamically displays the result (Disease Name, Confidence %).\n')
    p5.add_run('- It also renders the Recommended Treatments specific to the identified disease for the farmer to take action.')

    doc.add_page_break()
    doc.save(output_path)
    print(f"Document successfully created at {output_path}")

if __name__ == '__main__':
    img_path = sys.argv[1]
    out_path = sys.argv[2]
    create_doc(img_path, out_path)
