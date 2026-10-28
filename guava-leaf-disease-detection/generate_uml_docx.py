import sys
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_uml_doc(use_case_img, class_img, activity_img, output_path):
    doc = Document()
    
    # Title Page
    title = doc.add_heading('Guava Leaf Disease Detection', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = doc.add_paragraph('UML System Architecture Diagrams')
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].bold = True
    subtitle.runs[0].font.size = Pt(16)
    doc.add_page_break()

    # Use Case Diagram
    h1 = doc.add_heading('1. UML Use Case Diagram', level=1)
    doc.add_paragraph('This diagram illustrates the interactions between the primary actors (User/Farmer, Admin) and the system use cases such as Uploading Images, Running Inference, and Managing Datasets.')
    try:
        doc.add_picture(use_case_img, width=Inches(6.0))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    except Exception as e:
        print(f"Error loading Use Case Image: {e}")
    doc.add_page_break()

    # Class Diagram
    h2 = doc.add_heading('2. UML Class Diagram', level=1)
    doc.add_paragraph('This diagram represents the static structure of the application, detailing the Python classes such as StreamlitApp, ModelLoader, Preprocessor, InferenceEngine, and GradCamGenerator with their attributes and relationships.')
    try:
        doc.add_picture(class_img, width=Inches(6.0))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    except Exception as e:
        print(f"Error loading Class Image: {e}")
    doc.add_page_break()

    # Activity Diagram
    h3 = doc.add_heading('3. UML Activity Diagram', level=1)
    doc.add_paragraph('This diagram maps out the dynamic workflow of the image prediction process, from the start node and image upload through preprocessing, model inference, and displaying the final Grad-CAM result.')
    try:
        doc.add_picture(activity_img, width=Inches(6.0))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    except Exception as e:
        print(f"Error loading Activity Image: {e}")
    
    doc.save(output_path)
    print(f"UML Document successfully created at {output_path}")

if __name__ == '__main__':
    use_case = sys.argv[1]
    class_img = sys.argv[2]
    activity_img = sys.argv[3]
    out_path = sys.argv[4]
    create_uml_doc(use_case, class_img, activity_img, out_path)
