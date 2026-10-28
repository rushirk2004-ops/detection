import sys
import requests
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import time

diagrams = {
    "Use Case Diagram": """@startuml
skinparam handwritten false
skinparam monochrome false
skinparam defaultFontName Arial
left to right direction
actor "User/Farmer" as user
package "Guava Disease Detection Application" {
  usecase "Upload Image" as UC1
  usecase "Select Model" as UC2
  usecase "View Prediction" as UC3
  usecase "View Confidence" as UC4
  usecase "View Grad-CAM" as UC5
}
user --> UC1
user --> UC2
user --> UC3
user --> UC4
user --> UC5
@enduml""",

    "Class Diagram": """@startuml
skinparam defaultFontName Arial
class StreamlitApp {
  + main()
  + display_results()
}
class Preprocessor {
  + resize_image(size=224)
  + normalize_tensor()
}
class InferenceEngine {
  + load_model(path)
  + run_forward_pass()
}
class GradCamVisualizer {
  + generate_heatmap()
  + overlay_heatmap()
}

StreamlitApp --> Preprocessor : uses >
StreamlitApp --> InferenceEngine : uses >
StreamlitApp --> GradCamVisualizer : uses >
@enduml""",

    "Activity Diagram": """@startuml
skinparam defaultFontName Arial
start
:User Uploads Image;
:Preprocess Image (Resize/Normalize);
:Model Forward Pass;
:Apply Softmax;
:Generate Grad-CAM Map;
:Display Disease & Confidence;
stop
@enduml""",

    "Sequence Diagram": """@startuml
skinparam defaultFontName Arial
actor User
participant UI
participant Preprocessor
participant PyTorchModel
participant GradCAM

User -> UI : Upload Image
UI -> Preprocessor : Process Image
Preprocessor --> UI : Tensor
UI -> PyTorchModel : Forward Pass
PyTorchModel --> UI : Logits
UI -> GradCAM : Request Heatmap
GradCAM --> UI : Heatmap
UI -> User : Display Results
@enduml""",

    "Component Diagram": """@startuml
skinparam defaultFontName Arial
package "Application" {
  [User Interface]
  [Preprocessor]
  [Inference Engine]
  [Explainability]
}
node "Storage" {
  [Model Weights (.pt)]
}

[User Interface] --> [Preprocessor]
[Preprocessor] --> [Inference Engine]
[Inference Engine] --> [Model Weights (.pt)]
[Inference Engine] --> [Explainability]
@enduml""",

    "Deployment Diagram": """@startuml
skinparam defaultFontName Arial
node "Client Device" {
  [Web Browser]
}

node "Compute Node" {
  [Streamlit Server]
  [PyTorch Framework]
}
folder "Models" {
  artifact "efficientnet_best.pt"
  artifact "resnet50_best.pt"
}

[Web Browser] <--> [Streamlit Server]
[Streamlit Server] --> [PyTorch Framework]
[PyTorch Framework] --> "Models"
@enduml"""
}

def generate_image_from_kroki(name, puml_text):
    url = 'https://kroki.io/plantuml/png'
    headers = {'Content-Type': 'text/plain'}
    response = requests.post(url, headers=headers, data=puml_text.encode('utf-8'))
    if response.status_code == 200:
        filename = f"{name.replace(' ', '_').lower()}.png"
        with open(filename, 'wb') as f:
            f.write(response.content)
        return filename
    else:
        print(f"Failed to generate {name}. Status code: {response.status_code}")
        print(f"Response: {response.text}")
        return None

def create_doc(output_path):
    doc = Document()
    
    title = doc.add_heading('Guava Leaf & Fruit Disease Detection System', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = doc.add_paragraph('Comprehensive UML System Architecture Diagrams')
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].bold = True
    subtitle.runs[0].font.size = Pt(16)
    
    desc = doc.add_paragraph('This document contains the exact UML diagrams representing the real, self-contained architecture of the project (Streamlit, PyTorch, EfficientNet/ResNet, Grad-CAM). No unnecessary external components (like separate databases or REST APIs) have been included.')
    desc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    desc.runs[0].italic = True
    
    doc.add_page_break()

    for idx, (name, puml) in enumerate(diagrams.items(), 1):
        print(f"Generating image for {name}...")
        img_path = generate_image_from_kroki(name, puml)
        time.sleep(1)
        
        if img_path:
            h = doc.add_heading(f'{idx}. {name}', level=1)
            try:
                pic = doc.add_picture(img_path, width=Inches(6.0))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            except Exception as e:
                print(f"Error loading image {img_path}: {e}")
            
            if idx < len(diagrams):
                doc.add_page_break()
                
    doc.save(output_path)
    print(f"UML Document successfully created at {output_path}")

if __name__ == '__main__':
    out_path = sys.argv[1]
    create_doc(out_path)
