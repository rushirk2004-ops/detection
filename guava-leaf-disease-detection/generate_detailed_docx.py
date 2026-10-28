import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE

def add_flow_block(doc, text):
    lines = text.strip().split('\n')
    for line in lines:
        p = doc.add_paragraph(line.strip())
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].font.size = Pt(10)
        if '↓' in line or '▼' in line:
            p.runs[0].font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        else:
            p.runs[0].bold = True

def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    for i, hdr in enumerate(headers):
        hdr_cells[i].text = hdr
        hdr_cells[i].paragraphs[0].runs[0].bold = True
        
    for row_data in rows:
        row_cells = table.add_row().cells
        for i, val in enumerate(row_data):
            row_cells[i].text = str(val)
    doc.add_paragraph()

def create_detailed_doc(leaf_img, fruit_img, flow_img, output_path):
    doc = Document()
    
    # Title
    title = doc.add_heading('GUAVA LEAF & FRUIT DISEASE DETECTION SYSTEM', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = doc.add_paragraph('Real System Architecture and End-to-End Technical Flow\nPyTorch + EfficientNet-B3 / ResNet-50 + Streamlit + Grad-CAM')
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].bold = True
    subtitle.runs[0].font.size = Pt(14)
    
    doc.add_paragraph('This document describes the comprehensive system architecture for both Guava Leaf and Guava Fruit disease detection applications. It outlines the project structure, user interface, model loading, preprocessing, prediction, confidence calculation, and Grad-CAM explainability paths.').italic = True

    # Section: Project Layer
    doc.add_heading('System Technology Stack', level=1)
    stack_data = [
        ("User Interface", "Streamlit"),
        ("Image Handling", "PIL, NumPy, OpenCV"),
        ("Deep Learning", "PyTorch + torchvision"),
        ("Models", "EfficientNet-B3 and ResNet-50"),
        ("Explainable AI", "Grad-CAM"),
        ("Input Size", "224 × 224"),
        ("Supported Input", "JPG, JPEG, PNG"),
        ("Execution Device", "CUDA GPU when available; otherwise CPU")
    ]
    add_table(doc, ["Technology / Component", "Details"], stack_data)

    # 1. System Architecture Definition
    doc.add_heading('1. System Architecture Overview', level=1)
    doc.add_paragraph('System architecture is the high-level structure of the complete software system. It explains which major components exist, what each component does, how components communicate, and how data moves from the user\'s input to the final output.')
    
    # Part I: Leaf Architecture
    doc.add_heading('PART I: GUAVA LEAF DISEASE DETECTION', level=1)
    try:
        pic = doc.add_picture(leaf_img, width=Inches(6.0))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap = doc.add_paragraph('Figure 1. Complete architecture of the Guava Leaf Disease Detection System')
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.runs[0].italic = True
    except Exception as e:
        print(f"Leaf image error: {e}")

    doc.add_heading('Leaf Disease Classes', level=2)
    add_table(doc, ["Class", "Meaning in classifier"], [
        ("canker", "Guava leaf disease class"),
        ("dot", "Guava leaf disease class"),
        ("healthy", "Healthy leaf reference class"),
        ("mummification", "Guava leaf disease class"),
        ("rust", "Guava leaf disease class")
    ])

    # Part II: Fruit Architecture
    doc.add_heading('PART II: GUAVA FRUIT DISEASE DETECTION', level=1)
    try:
        pic = doc.add_picture(fruit_img, width=Inches(6.0))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap = doc.add_paragraph('Figure 2. Complete architecture of the Guava Fruit Disease Detection System')
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.runs[0].italic = True
    except Exception as e:
        print(f"Fruit image error: {e}")

    # Core Pipeline
    doc.add_heading('CORE ARCHITECTURAL LAYERS & PIPELINE (Leaf & Fruit)', level=1)
    layers_data = [
        ("1", "Dataset Layer", "Stores images organized by disease/healthy class."),
        ("2", "Data Preparation Layer", "Merges, normalizes, splits, preprocesses, and augments images."),
        ("3", "Model Layer", "Provides EfficientNet-B3 and ResNet-50 classification models."),
        ("4", "Training Layer", "Trains models and saves best validation checkpoints."),
        ("5", "Model Storage Layer", "Stores trained .pt model artifacts."),
        ("6", "Streamlit Presentation Layer", "Provides upload, configuration, and result interface."),
        ("7", "Inference Layer", "Preprocesses uploaded image and produces class probabilities."),
        ("8", "Explainability Layer", "Uses Grad-CAM to visualize image regions influencing prediction."),
        ("9", "Output Layer", "Displays disease prediction, confidence, and Grad-CAM image.")
    ]
    add_table(doc, ["Layer", "Component", "Main Responsibility"], layers_data)

    # NEW SECTION: SYSTEM EXECUTION FLOWCHART
    doc.add_heading('SYSTEM EXECUTION FLOWCHART', level=1)
    doc.add_paragraph('The following flowchart illustrates the complex end-to-end technical sequence of execution in a highly readable format.')
    
    try:
        pic = doc.add_picture(flow_img, width=Inches(6.0))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap = doc.add_paragraph('Figure 3. Step-by-step execution flowchart of the System')
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.runs[0].italic = True
    except Exception as e:
        print(f"Flow image error: {e}")

    doc.add_heading('Step-by-Step Flowchart Explanation', level=2)
    flow_steps = [
        "Step 1: Streamlit UI Initializes - The system sets the page configuration, UI layout, and detects if CUDA (GPU) or CPU is available for execution.",
        "Step 2: Load PyTorch Model (.pt checkpoint) - Based on user selection (EfficientNet or ResNet), the system constructs the model and loads the pre-trained weights into memory, setting model.eval().",
        "Step 3: User Uploads Image - The user selects a JPG or PNG image of a Guava leaf or fruit through the Streamlit uploader.",
        "Step 4: Image Preprocessing - The PIL image is converted to RGB, resized to exactly 224x224, converted to a PyTorch Tensor, and normalized using standard ImageNet statistics.",
        "Step 5: GPU/CPU Inference - The preprocessed image tensor is passed into the Deep Learning model for a forward pass (under torch.no_grad()) to extract deep visual features.",
        "Step 6: Softmax Probabilities - The raw logits from the model are passed through a Softmax function to convert them into a percentage-based probability distribution. The class with the highest probability is selected (Argmax).",
        "Step 7: Grad-CAM Explainability - The system selects the target Convolutional Layer (e.g., layer4 for ResNet) and generates a Gradient-weighted Class Activation Mapping (Grad-CAM) heatmap to explain the model's focus.",
        "Step 8: Display Final Results - The Streamlit UI displays the Predicted Disease Class, the Confidence Percentage, and overlays the Grad-CAM heatmap on the original image for the user to interpret."
    ]
    for step in flow_steps:
        p = doc.add_paragraph(step, style='List Bullet')
        
    doc.add_page_break()
    doc.save(output_path)
    print(f"Final Document with Flowchart successfully created at {output_path}")

if __name__ == '__main__':
    leaf_img_path = sys.argv[1]
    fruit_img_path = sys.argv[2]
    flow_img_path = sys.argv[3]
    out_path = sys.argv[4]
    create_detailed_doc(leaf_img_path, fruit_img_path, flow_img_path, out_path)
